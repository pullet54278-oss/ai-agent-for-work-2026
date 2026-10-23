#!/usr/bin/env python3
"""Recalculate sample-workbook.xlsx in LibreOffice (headless) and compare every formula cell with the
independently computed values from make_workbook.py. Writes docs/excel-workbook-report.json."""
import json, os, shutil, subprocess, sys, tempfile
from datetime import datetime, date
from openpyxl import load_workbook

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
SRC = os.path.join(ROOT, 'ai-excel-formula-mastery', 'data', 'sample-workbook.xlsx')
EXP = json.load(open(os.path.join(ROOT, 'scripts', 'excel', '.workbook-expected.json'), encoding='utf-8'))
REQUIRED = ['Raw_Data', 'Master_Data', 'Calculations', 'Validation', 'Dashboard', 'Instructions']

tmp = tempfile.mkdtemp(prefix='wbverify-')
try:
    subprocess.run(['soffice', f'-env:UserInstallation=file://{tmp}/profile', '--headless', '--calc', '--convert-to', 'xlsx', '--outdir', tmp, SRC],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    wb = load_workbook(os.path.join(tmp, os.path.basename(SRC)), data_only=True)
    src_wb = load_workbook(SRC)
    fails, ok = [], 0
    for ref, want in EXP.items():
        sh, cell = ref.split('!')
        got = wb[sh][cell].value
        if isinstance(got, datetime): got = (got - datetime(1899, 12, 30)).days + (got - datetime(1899, 12, 30)).seconds / 86400
        elif isinstance(got, date): got = (got - date(1899, 12, 30)).days
        good = (abs(got - want) <= 1e-6 * max(1, abs(want))) if isinstance(want, (int, float)) and isinstance(got, (int, float)) else (str(got) == str(want))
        if good: ok += 1
        else: fails.append({'cell': ref, 'formula': src_wb[sh][cell].value, 'want': want, 'got': got})
    sheets = wb.sheetnames
    missing = [s for s in REQUIRED if s not in sheets]
    report = {'generatedAt': datetime.now().isoformat(timespec='seconds'), 'engine': 'LibreOffice headless recalculation', 'file': os.path.relpath(SRC, ROOT),
              'sheets': sheets, 'requiredSheetsMissing': missing, 'formulaCells': len(EXP), 'passed': ok, 'failed': len(fails), 'fails': fails,
              'rows': {s: wb[s].max_row for s in sheets}}
    json.dump(report, open(os.path.join(ROOT, 'docs', 'excel-workbook-report.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'workbook: {len(sheets)} sheets (missing required: {missing or "none"}) · formula cells {ok}/{len(EXP)} match')
    for x in fails[:30]: print('  ✗', x['cell'], x['formula'], 'want', x['want'], 'got', x['got'])
    sys.exit(1 if fails or missing else 0)
finally:
    shutil.rmtree(tmp, ignore_errors=True)
