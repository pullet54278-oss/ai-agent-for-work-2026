#!/usr/bin/env python3
"""Build ai-excel-formula-mastery/data/sample-workbook.xlsx (+ content/excel/sample-data.json for the web dashboard).

Fictional practice data (ไม่ใช่มาตรฐานสายพันธุ์). Every formula cell also gets an independently computed
expected value in `EXPECTED`, written to scripts/excel/.workbook-expected.json, which verify_workbook.py
compares against a real LibreOffice recalculation.

KPI definitions used everywhere (stated in the Instructions sheet):
  Live_Birds      = Birds_Placed - Mortality_Cum                       (ตัว)
  Weight_Gain_kg  = (Body_Weight_g - Chick_Weight_g) / 1000 * Live_Birds (กก. น้ำหนักที่เพิ่มของไก่ที่ยังมีชีวิต)
  FCR             = Feed_Cum_kg / Weight_Gain_kg                        (อาหารสะสมตั้งแต่ลงไก่ ÷ น้ำหนักที่เพิ่ม)
  Mortality_%     = Mortality_Cum / Birds_Placed * 100                  (สะสม)
These are teaching definitions — confirm the organisation's standard before real use.
"""
import json, math, os, statistics, sys
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import FormulaRule

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'ai-excel-formula-mastery', 'data', 'sample-workbook.xlsx')
CHICK_G = 42
REPORT_DATE = date(2026, 10, 31)
EXPECTED = {}  # 'Sheet!A1' -> value

# ---------------- data (deterministic) ----------------
FARMS = [  # Farm_ID, Region, Houses, Birds_Placed, Placement_Date, Target_FCR_42d, bw factor, feed factor, mort/wk
    ('F001', 'ภาคกลาง', 4, 10000, date(2026, 9, 17), 1.65, 1.00, 1.00, 0.0040),
    ('F002', 'ภาคตะวันออก', 4, 9800, date(2026, 9, 17), 1.65, 0.97, 1.01, 0.0055),
    ('F003', 'ภาคเหนือ', 6, 15000, date(2026, 9, 17), 1.68, 1.02, 0.99, 0.0035),
    ('F004', 'ภาคอีสาน', 5, 12000, date(2026, 9, 17), 1.68, 0.95, 1.04, 0.0070),
]
AGES = [7, 14, 21, 28, 35, 42]
BW_CURVE = {7: 185, 14: 510, 21: 940, 28: 1480, 35: 2080, 42: 2690}       # g (สมมติ)
FCR_CURVE = {7: 0.95, 14: 1.12, 21: 1.30, 28: 1.42, 35: 1.53, 42: 1.64}   # สมมติ
ROWS = []
for fid, region, houses, placed, pdate, tfcr, bwf, ff, mwk in FARMS:
    for i, age in enumerate(AGES):
        bw = round(BW_CURVE[age] * bwf)
        mort = round(placed * mwk * (i + 1) + (3 if fid == 'F004' and age >= 28 else 0))
        live = placed - mort
        gain = (bw - CHICK_G) / 1000 * live
        feed = round(FCR_CURVE[age] * ff * gain)
        ROWS.append({'Farm_ID': fid, 'Date': pdate + timedelta(days=age), 'Age_Days': age, 'Body_Weight_g': bw,
                     'Feed_Cum_kg': feed, 'Mortality_Cum': mort, 'Birds_Placed': placed})

SPEC_SAMPLE = [  # spec section 15, verbatim
    ('F001', date(2026, 10, 1), 14, 520, 2400, 3, 10000), ('F001', date(2026, 10, 8), 21, 950, 3600, 4, 10000),
    ('F002', date(2026, 10, 1), 14, 505, 2350, 2, 9800), ('F002', date(2026, 10, 8), 21, 925, 3500, 5, 9800)]
SALES = [('O001', date(2026, 10, 1), 'P001', 'Online', 12, 150), ('O002', date(2026, 10, 2), 'P002', 'Retail', 8, 220),
         ('O003', date(2026, 10, 3), 'P001', 'Online', 10, 150), ('O004', date(2026, 10, 3), 'P003', 'Retail', 5, 480),
         ('O005', date(2026, 10, 4), 'P002', 'Online', 3, 220), ('O006', date(2026, 10, 5), 'P009', 'Online', 2, 99)]
PRODUCTS = [('P001', 'Product A', 'Feed', 150), ('P002', 'Product B', 'Equipment', 220), ('P003', 'Product C', 'Equipment', 480)]
# 100 bird weights (g) at day 21 for F001 — fixed list, deterministic
WEIGHTS = [940 + round(95 * math.sin(i * 1.7) + 60 * math.cos(i * 0.9)) for i in range(100)]

# ---------------- styles ----------------
HEAD = PatternFill('solid', fgColor='2557E0'); HEAD_F = Font(bold=True, color='FFFFFF')
NOTE_F = Font(italic=True, color='5A6275'); TITLE_F = Font(bold=True, size=14, color='1A2233')
INPUT = PatternFill('solid', fgColor='FFF4CE'); CALC = PatternFill('solid', fgColor='E8F0FE')
thin = Side(style='thin', color='D0D5E0'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)


def header(ws, row, cols):
    for c, name in enumerate(cols, 1):
        cell = ws.cell(row=row, column=c, value=name); cell.fill = HEAD; cell.font = HEAD_F; cell.border = BOX
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)


def widths(ws, w):
    for i, x in enumerate(w, 1): ws.column_dimensions[get_column_letter(i)].width = x


def f(ws, ref, formula, expected, fmt=None, fill=CALC):
    """write a formula cell and remember its independently computed expected value"""
    c = ws[ref]; c.value = formula; c.fill = fill; c.border = BOX
    if fmt: c.number_format = fmt
    EXPECTED[f'{ws.title}!{ref}'] = expected


wb = Workbook()
# ---------------- Instructions ----------------
ins = wb.active; ins.title = 'Instructions'
lines = [
    ('AI + Excel Formula Mastery — Sample Workbook', TITLE_F),
    ('ข้อมูลทั้งหมดเป็นข้อมูลสมมติเพื่อฝึก ไม่ใช่มาตรฐานสายพันธุ์และไม่ใช่ข้อมูลของหน่วยงานใด', NOTE_F),
    ('', None),
    ('ชีตในไฟล์', Font(bold=True)),
    ('Raw_Data — ข้อมูลผลการเลี้ยงรายสัปดาห์ 4 ฟาร์ม (F001–F004) อายุ 7–42 วัน (Excel Table: tblRaw)', None),
    ('Master_Data — ข้อมูลหลักของฟาร์ม + Product Master (ใช้ฝึก Lookup)', None),
    ('Calculations — KPI ต่อแถว (Live Birds, Weight Gain, FCR, Mortality %) + สรุปรายฟาร์ม + สถิติน้ำหนัก (%CV)', None),
    ('Validation — ตัวอย่างข้อมูลที่มีปัญหา + สูตร Data Quality Status (PASS/WARNING/ERROR) พร้อมเหตุผล', None),
    ('Dashboard — เลือกฟาร์มในเซลล์ C3 แล้ว KPI และกราฟจะเปลี่ยนตาม', None),
    ('Sales_Data / Product_Master — ข้อมูลยอดขายจากต้นฉบับ (ข้อ 15) + แถวเพิ่มสำหรับทดสอบ', None),
    ('Farm_Performance — ตาราง Farm_Performance จากต้นฉบับข้อ 15 ตามเดิม (ดูข้อจำกัดในชีตนั้น)', None),
    ('Weight_Samples — น้ำหนักไก่ตัวอย่าง 100 ตัว (กรัม) สำหรับฝึก %CV / Percentile', None),
    ('M365_Formulas — สูตร Microsoft 365 (XLOOKUP, FILTER, LET, LAMBDA ฯลฯ) เป็นข้อความให้คัดลอกไปลอง พร้อมสูตรทางเลือก', None),
    ('', None),
    ('นิยาม KPI (ใช้เพื่อการเรียน — ยืนยันนิยามกับมาตรฐานองค์กรก่อนใช้งานจริง)', Font(bold=True)),
    ('Live_Birds = Birds_Placed − Mortality_Cum (ตัว)', None),
    (f'Weight_Gain_kg = (Body_Weight_g − {CHICK_G}) ÷ 1000 × Live_Birds  — สมมติน้ำหนักลูกไก่แรกรับ {CHICK_G} กรัม (ปรับได้ที่ Calculations!N2)', None),
    ('FCR = Feed_Cum_kg ÷ Weight_Gain_kg  — Feed_Cum_kg คืออาหารสะสมตั้งแต่ลงไก่ถึงวันที่บันทึก (ช่วงเวลาเดียวกับน้ำหนัก)', None),
    ('Mortality_% = Mortality_Cum ÷ Birds_Placed × 100 (สะสม)', None),
    ('', None),
    ('วิธีตรวจสอบ', Font(bold=True)),
    ('1) เซลล์สีเหลือง = ข้อมูลนำเข้า, สีฟ้า = สูตร · 2) ลองแก้ค่าใน Raw_Data แล้วดู Calculations/Dashboard เปลี่ยน', None),
    ('3) ทุกสูตรในไฟล์นี้ตรวจผลด้วยการคำนวณใหม่ใน LibreOffice เทียบกับค่าที่คำนวณแยกด้วย Python (ดูรายงาน docs/excel-workbook-report.json)', None),
    (f'4) วันที่รายงาน (Report Date) ที่ใช้ตรวจ "วันที่เกินวันปัจจุบัน" อยู่ที่เซลล์ B27 = {REPORT_DATE.isoformat()} (เปลี่ยนเป็น =TODAY() ได้)', None),
    ('', None),
    ('Report Date', Font(bold=True)),
]
for i, (t, font) in enumerate(lines, 1):
    c = ins.cell(row=i, column=1, value=t)
    if font: c.font = font
ins['B27'] = REPORT_DATE; ins['B27'].number_format = 'yyyy-mm-dd'; ins['B27'].fill = INPUT
ins['A28'] = 'เวอร์ชันที่รองรับ: สูตรในชีตหลักใช้ฟังก์ชันที่มีใน Excel 2010 ขึ้นไป (MAXIFS ต้อง 2019+) · สูตร Microsoft 365 อยู่ในชีต M365_Formulas'
widths(ins, [120, 16])

# ---------------- Raw_Data ----------------
raw = wb.create_sheet('Raw_Data')
cols = ['Farm_ID', 'Date', 'Age_Days', 'Body_Weight_g', 'Feed_Cum_kg', 'Mortality_Cum', 'Birds_Placed']
header(raw, 1, cols)
for r, row in enumerate(ROWS, 2):
    for c, k in enumerate(cols, 1):
        cell = raw.cell(row=r, column=c, value=row[k]); cell.fill = INPUT; cell.border = BOX
        if k == 'Date': cell.number_format = 'yyyy-mm-dd'
        if k in ('Feed_Cum_kg', 'Birds_Placed'): cell.number_format = '#,##0'
NRAW = len(ROWS); LAST = NRAW + 1
raw.add_table(Table(displayName='tblRaw', ref=f'A1:G{LAST}', tableStyleInfo=TableStyleInfo(name='TableStyleMedium2', showRowStripes=True)))
widths(raw, [10, 12, 10, 15, 14, 15, 13]); raw.freeze_panes = 'A2'

# ---------------- Master_Data ----------------
ms = wb.create_sheet('Master_Data')
header(ms, 1, ['Farm_ID', 'Region', 'Houses', 'Birds_Placed', 'Placement_Date', 'Target_FCR_42d'])
for r, (fid, region, houses, placed, pdate, tfcr, *_r) in enumerate(FARMS, 2):
    for c, v in enumerate([fid, region, houses, placed, pdate, tfcr], 1):
        cell = ms.cell(row=r, column=c, value=v); cell.fill = INPUT; cell.border = BOX
        if c == 5: cell.number_format = 'yyyy-mm-dd'
header(ms, 8, ['Product_ID', 'Product_Name', 'Category', 'Price'])
for r, p in enumerate(PRODUCTS, 9):
    for c, v in enumerate(p, 1): cell = ms.cell(row=r, column=c, value=v); cell.fill = INPUT; cell.border = BOX
widths(ms, [11, 16, 11, 13, 15, 15])

# ---------------- Calculations ----------------
ca = wb.create_sheet('Calculations')
header(ca, 1, ['Farm_ID', 'Date', 'Age_Days', 'Live_Birds', 'Weight_Gain_kg', 'FCR', 'Mortality_%', 'Target_FCR', 'FCR_Status', 'Week_Label'])
ca['M1'] = 'Chick_Weight_g'; ca['M1'].font = Font(bold=True); ca['N1'] = 'ค่าที่ปรับได้'
ca['M2'] = 'น้ำหนักลูกไก่แรกรับ (กรัม)'; ca['N2'] = CHICK_G; ca['N2'].fill = INPUT
tgt = {fid: tfcr for fid, _a, _b, _c, _d, tfcr, *_r in FARMS}
for i, row in enumerate(ROWS):
    r = i + 2
    live = row['Birds_Placed'] - row['Mortality_Cum']
    gain = (row['Body_Weight_g'] - CHICK_G) / 1000 * live
    fcr = row['Feed_Cum_kg'] / gain
    mort = row['Mortality_Cum'] / row['Birds_Placed'] * 100
    f(ca, f'A{r}', f'=Raw_Data!A{r}', row['Farm_ID'])
    f(ca, f'B{r}', f'=Raw_Data!B{r}', (row['Date'] - date(1899, 12, 30)).days, 'yyyy-mm-dd')
    f(ca, f'C{r}', f'=Raw_Data!C{r}', row['Age_Days'])
    f(ca, f'D{r}', f'=Raw_Data!G{r}-Raw_Data!F{r}', live, '#,##0')
    f(ca, f'E{r}', f'=IF(AND(ISNUMBER(Raw_Data!D{r}),Raw_Data!D{r}>$N$2),(Raw_Data!D{r}-$N$2)/1000*D{r},NA())', gain, '#,##0.0')
    f(ca, f'F{r}', f'=IFERROR(IF(E{r}>0,Raw_Data!E{r}/E{r},NA()),NA())', fcr, '0.000')
    f(ca, f'G{r}', f'=IF(Raw_Data!G{r}>0,Raw_Data!F{r}/Raw_Data!G{r}*100,NA())', mort, '0.00')
    f(ca, f'H{r}', f'=IFERROR(INDEX(Master_Data!$F$2:$F$5,MATCH(A{r},Master_Data!$A$2:$A$5,0)),"ไม่พบฟาร์ม")', tgt[row['Farm_ID']], '0.00')
    status = 'สูงกว่าเป้า' if (row['Age_Days'] == 42 and fcr > tgt[row['Farm_ID']]) else ('ตามเป้า' if row['Age_Days'] == 42 else 'ยังไม่ถึง 42 วัน')
    f(ca, f'I{r}', f'=IF(C{r}<42,"ยังไม่ถึง 42 วัน",IF(F{r}>H{r},"สูงกว่าเป้า","ตามเป้า"))', status)
    f(ca, f'J{r}', f'="สัปดาห์ "&ROUNDUP(C{r}/7,0)&" · "&TEXT(B{r},"yyyy-mm-dd")', f"สัปดาห์ {math.ceil(row['Age_Days'] / 7)} · {row['Date'].isoformat()}")
# per-farm summary (latest age = 42)
S0 = NRAW + 4
ca.cell(row=S0 - 1, column=1, value='สรุปรายฟาร์ม (อายุล่าสุด)').font = Font(bold=True, size=12)
header(ca, S0, ['Farm_ID', 'Last_Age', 'Last_Date', 'Body_Weight_g', 'FCR', 'Mortality_%', 'Feed_Cum_kg', 'Rows', 'Avg_BW_all_rows'])
for k, (fid, *_r) in enumerate(FARMS):
    r = S0 + 1 + k
    rows = [x for x in ROWS if x['Farm_ID'] == fid]; last = max(rows, key=lambda x: x['Age_Days'])
    live = last['Birds_Placed'] - last['Mortality_Cum']; gain = (last['Body_Weight_g'] - CHICK_G) / 1000 * live
    ca[f'A{r}'] = fid; ca[f'A{r}'].fill = INPUT
    rng = lambda col: f'Raw_Data!${col}$2:${col}${LAST}'
    f(ca, f'B{r}', f'=_xlfn.MAXIFS({rng("C")},{rng("A")},A{r})', last['Age_Days'])
    f(ca, f'C{r}', f'=_xlfn.MAXIFS({rng("B")},{rng("A")},A{r})', (last['Date'] - date(1899, 12, 30)).days, 'yyyy-mm-dd')
    f(ca, f'D{r}', f'=SUMIFS({rng("D")},{rng("A")},A{r},{rng("C")},B{r})', last['Body_Weight_g'])
    f(ca, f'E{r}', f'=SUMIFS($F$2:$F${LAST},$A$2:$A${LAST},A{r},$C$2:$C${LAST},B{r})', last['Feed_Cum_kg'] / gain, '0.000')
    f(ca, f'F{r}', f'=SUMIFS($G$2:$G${LAST},$A$2:$A${LAST},A{r},$C$2:$C${LAST},B{r})', last['Mortality_Cum'] / last['Birds_Placed'] * 100, '0.00')
    f(ca, f'G{r}', f'=SUMIFS({rng("E")},{rng("A")},A{r},{rng("C")},B{r})', last['Feed_Cum_kg'], '#,##0')
    f(ca, f'H{r}', f'=COUNTIF({rng("A")},A{r})', len(rows))
    f(ca, f'I{r}', f'=AVERAGEIFS({rng("D")},{rng("A")},A{r})', sum(x['Body_Weight_g'] for x in rows) / len(rows), '0.0')
# weight statistics (I03 / A01)
W0 = S0 + 7
ca.cell(row=W0, column=1, value='สถิติน้ำหนักตัวอย่าง F001 อายุ 21 วัน (Weight_Samples!B2:B101, กรัม)').font = Font(bold=True, size=12)
ws_rng = 'Weight_Samples!$B$2:$B$101'
st = {
    'n': len(WEIGHTS), 'Mean': statistics.mean(WEIGHTS), 'Median': statistics.median(WEIGHTS), 'SD (Sample)': statistics.stdev(WEIGHTS),
}
st['CV%'] = st['SD (Sample)'] / st['Mean'] * 100
def pinc(data, p):  # Excel PERCENTILE.INC
    s = sorted(data); h = (len(s) - 1) * p; lo = math.floor(h); return s[lo] + (h - lo) * (s[min(lo + 1, len(s) - 1)] - s[lo])
st['P10'] = pinc(WEIGHTS, 0.10); st['P90'] = pinc(WEIGHTS, 0.90); st['IQR'] = pinc(WEIGHTS, 0.75) - pinc(WEIGHTS, 0.25)
forms = {
    'n': f'=COUNT({ws_rng})', 'Mean': f'=AVERAGE({ws_rng})', 'Median': f'=MEDIAN({ws_rng})', 'SD (Sample)': f'=_xlfn.STDEV.S({ws_rng})',
    'CV%': f'=IF(OR(COUNT({ws_rng})<2,AVERAGE({ws_rng})=0),NA(),_xlfn.STDEV.S({ws_rng})/AVERAGE({ws_rng})*100)',
    'P10': f'=_xlfn.PERCENTILE.INC({ws_rng},0.1)', 'P90': f'=_xlfn.PERCENTILE.INC({ws_rng},0.9)',
    'IQR': f'=_xlfn.QUARTILE.INC({ws_rng},3)-_xlfn.QUARTILE.INC({ws_rng},1)',
}
for k, (name, formula) in enumerate(forms.items()):
    r = W0 + 1 + k
    ca[f'A{r}'] = name; ca[f'A{r}'].font = Font(bold=True)
    f(ca, f'B{r}', formula, st[name], '0' if name == 'n' else '0.00')
ca[f'C{W0 + 5}'] = 'อย่าตีความ %CV เป็นเกณฑ์มาตรฐานสายพันธุ์โดยไม่มีข้อมูลอ้างอิง'; ca[f'C{W0 + 5}'].font = NOTE_F
widths(ca, [11, 12, 10, 12, 15, 9, 12, 11, 17, 26, 3, 3, 26, 14]); ca.freeze_panes = 'A2'

# ---------------- Weight_Samples ----------------
wsm = wb.create_sheet('Weight_Samples')
header(wsm, 1, ['Bird_No', 'Body_Weight_g'])
for i, w in enumerate(WEIGHTS, 2):
    wsm.cell(row=i, column=1, value=i - 1); c = wsm.cell(row=i, column=2, value=w); c.fill = INPUT
widths(wsm, [10, 15])

# ---------------- Validation (I06) ----------------
va = wb.create_sheet('Validation')
va['A1'] = 'Data Quality Audit — ข้อมูลตัวอย่างที่มีปัญหาโดยตั้งใจ (แก้ค่าแล้วดูสถานะเปลี่ยน)'; va['A1'].font = TITLE_F
header(va, 3, ['Farm_ID', 'Date', 'Body_Weight_g', 'Feed_Cum_kg', 'Mortality_Cum', 'Status', 'Reason'])
VROWS = [
    ('F001', date(2026, 9, 24), 185, 1595, 40),      # ok
    ('F001', date(2026, 10, 1), 510, 4400, 80),      # ok
    ('F002', date(2026, 10, 1), 495, 4321, 108),     # ok
    ('F002', date(2026, 10, 1), 495, 4321, 108),     # duplicate Farm+Date
    ('', date(2026, 10, 8), 940, 9000, 120),         # blank Farm_ID
    ('F003', date(2026, 10, 8), -950, 14000, 150),   # negative weight
    ('F003', date(2026, 11, 20), 1500, 20000, 190),  # date after report date
    ('F004', date(2026, 10, 8), '1,000', 11000, 200),  # number stored as text
    ('F004', date(2026, 10, 15), 1400, None, 260),   # blank feed (warning)
]
for i, (fid, d, bw, feed, mort) in enumerate(VROWS):
    r = 4 + i
    for c, v in enumerate([fid, d, bw, feed, mort], 1):
        cell = va.cell(row=r, column=c, value=v if v != '' else None); cell.fill = INPUT; cell.border = BOX
        if c == 2: cell.number_format = 'yyyy-mm-dd'
    # expected (independent python logic of the same rules)
    dup = sum(1 for x in VROWS if x[0] == fid and x[1] == d) > 1
    if not fid: exp = ('ERROR', 'ไม่มี Farm_ID')
    elif not isinstance(bw, (int, float)): exp = ('ERROR', 'Body_Weight_g ไม่ใช่ตัวเลข')
    elif bw < 0: exp = ('ERROR', 'ค่าติดลบ')
    elif d > REPORT_DATE: exp = ('ERROR', 'วันที่เกินวันรายงาน')
    elif dup: exp = ('WARNING', 'ข้อมูลซ้ำ (Farm_ID + Date)')
    elif feed is None: exp = ('WARNING', 'Feed_Cum_kg ว่าง')
    else: exp = ('PASS', 'ผ่าน')
    lastv = 4 + len(VROWS) - 1
    f(va, f'F{r}', (f'=IF(A{r}="","ERROR",IF(NOT(ISNUMBER(C{r})),"ERROR",IF(C{r}<0,"ERROR",IF(B{r}>Instructions!$B$27,"ERROR",'
                    f'IF(COUNTIFS($A$4:$A${lastv},A{r},$B$4:$B${lastv},B{r})>1,"WARNING",IF(D{r}="","WARNING","PASS"))))))'), exp[0])
    f(va, f'G{r}', (f'=IF(A{r}="","ไม่มี Farm_ID",IF(NOT(ISNUMBER(C{r})),"Body_Weight_g ไม่ใช่ตัวเลข",IF(C{r}<0,"ค่าติดลบ",IF(B{r}>Instructions!$B$27,"วันที่เกินวันรายงาน",'
                    f'IF(COUNTIFS($A$4:$A${lastv},A{r},$B$4:$B${lastv},B{r})>1,"ข้อมูลซ้ำ (Farm_ID + Date)",IF(D{r}="","Feed_Cum_kg ว่าง","ผ่าน"))))))'), exp[1])
n_err = sum(1 for i in range(len(VROWS)) if EXPECTED[f'Validation!F{4 + i}'] == 'ERROR')
n_warn = sum(1 for i in range(len(VROWS)) if EXPECTED[f'Validation!F{4 + i}'] == 'WARNING')
lastv = 3 + len(VROWS)
va['I3'] = 'สรุป'; va['I3'].font = Font(bold=True)
for k, (lab, val, exp) in enumerate([('ERROR', 'ERROR', n_err), ('WARNING', 'WARNING', n_warn), ('PASS', 'PASS', len(VROWS) - n_err - n_warn)]):
    va[f'I{4 + k}'] = lab; f(va, f'J{4 + k}', f'=COUNTIF($F$4:$F${lastv},"{val}")', exp)
va.conditional_formatting.add(f'F4:F{lastv}', FormulaRule(formula=['F4="ERROR"'], fill=PatternFill('solid', fgColor='FDE7E9')))
va.conditional_formatting.add(f'F4:F{lastv}', FormulaRule(formula=['F4="WARNING"'], fill=PatternFill('solid', fgColor='FFF4CE')))
va['A15'] = 'กฎตรวจรูปแบบ (format): ค่าว่าง, ชนิดข้อมูล, ค่าติดลบ, วันที่เกินวันรายงาน · กฎธุรกิจ (business): ข้อมูลซ้ำ Farm_ID+Date, Feed ว่าง = WARNING ให้ผู้รับผิดชอบตรวจ'
va['A15'].font = NOTE_F
widths(va, [10, 12, 15, 13, 15, 11, 30, 3, 10, 8])

# ---------------- Dashboard ----------------
db = wb.create_sheet('Dashboard')
db['A1'] = 'Farm Performance Dashboard (Prototype — ข้อมูลตัวอย่างในไฟล์นี้)'; db['A1'].font = TITLE_F
db['A3'] = 'เลือกฟาร์ม'; db['A3'].font = Font(bold=True)
db['C3'] = 'F001'; db['C3'].fill = INPUT
dv = DataValidation(type='list', formula1='=Master_Data!$A$2:$A$5', allow_blank=False); db.add_data_validation(dv); dv.add('C3')
sel = [x for x in ROWS if x['Farm_ID'] == 'F001']; last = max(sel, key=lambda x: x['Age_Days'])
live = last['Birds_Placed'] - last['Mortality_Cum']; gain = (last['Body_Weight_g'] - CHICK_G) / 1000 * live
kpis = [
    ('อายุล่าสุด (วัน)', f'=_xlfn.MAXIFS(Raw_Data!$C$2:$C${LAST},Raw_Data!$A$2:$A${LAST},$C$3)', last['Age_Days'], '0'),
    ('Body Weight (g)', f'=SUMIFS(Raw_Data!$D$2:$D${LAST},Raw_Data!$A$2:$A${LAST},$C$3,Raw_Data!$C$2:$C${LAST},$C$5)', last['Body_Weight_g'], '#,##0'),
    ('FCR', f'=SUMIFS(Calculations!$F$2:$F${LAST},Calculations!$A$2:$A${LAST},$C$3,Calculations!$C$2:$C${LAST},$C$5)', last['Feed_Cum_kg'] / gain, '0.000'),
    ('Mortality %', f'=SUMIFS(Calculations!$G$2:$G${LAST},Calculations!$A$2:$A${LAST},$C$3,Calculations!$C$2:$C${LAST},$C$5)', last['Mortality_Cum'] / last['Birds_Placed'] * 100, '0.00'),
    ('Feed Intake สะสม (kg)', f'=SUMIFS(Raw_Data!$E$2:$E${LAST},Raw_Data!$A$2:$A${LAST},$C$3,Raw_Data!$C$2:$C${LAST},$C$5)', last['Feed_Cum_kg'], '#,##0'),
    ('Target FCR', f'=IFERROR(INDEX(Master_Data!$F$2:$F$5,MATCH($C$3,Master_Data!$A$2:$A$5,0)),"ไม่พบฟาร์ม")', tgt['F001'], '0.00'),
]
for k, (lab, formula, exp, fmt) in enumerate(kpis):
    r = 5 + k
    db[f'A{r}'] = lab; db[f'A{r}'].font = Font(bold=True)
    f(db, f'C{r}', formula, exp, fmt)
db['A12'] = 'Data Source'; db['C12'] = 'Raw_Data (ในไฟล์นี้) — ไม่เชื่อมแหล่งข้อมูลภายนอก'
db['A13'] = 'Last Updated'; f(db, 'C13', f'=_xlfn.MAXIFS(Raw_Data!$B$2:$B${LAST},Raw_Data!$A$2:$A${LAST},"<>")', (max(x['Date'] for x in ROWS) - date(1899, 12, 30)).days, 'yyyy-mm-dd')
db['A14'] = 'Data Quality'; f(db, 'C14', '=IF(Validation!$J$4>0,"มี "&Validation!$J$4&" แถว ERROR ในชีต Validation","ไม่มี ERROR")', f'มี {n_err} แถว ERROR ในชีต Validation')
# chart table: body weight by age for the 4 farms
db['E4'] = 'Age_Days'
for c, (fid, *_r) in enumerate(FARMS): db.cell(row=4, column=6 + c, value=fid).font = Font(bold=True)
for i, age in enumerate(AGES):
    r = 5 + i; db[f'E{r}'] = age
    for c, (fid, *_r) in enumerate(FARMS):
        col = get_column_letter(6 + c)
        exp = next(x['Body_Weight_g'] for x in ROWS if x['Farm_ID'] == fid and x['Age_Days'] == age)
        f(db, f'{col}{r}', f'=SUMIFS(Raw_Data!$D$2:$D${LAST},Raw_Data!$A$2:$A${LAST},{col}$4,Raw_Data!$C$2:$C${LAST},$E{r})', exp, '#,##0')
ch = LineChart(); ch.title = 'Body Weight (g) ตามอายุ'; ch.y_axis.title = 'กรัม'; ch.x_axis.title = 'อายุ (วัน)'; ch.height = 8; ch.width = 16
ch.add_data(Reference(db, min_col=6, max_col=9, min_row=4, max_row=10), titles_from_data=True)
ch.set_categories(Reference(db, min_col=5, min_row=5, max_row=10)); db.add_chart(ch, 'E13')
widths(db, [22, 2, 30, 2, 10, 10, 10, 10, 10])

# ---------------- Sales_Data / Product_Master (spec section 15 + extra test rows) ----------------
sd = wb.create_sheet('Sales_Data')
header(sd, 1, ['Order_ID', 'Order_Date', 'Product_ID', 'Channel', 'Quantity', 'Unit_Price', 'Amount', 'Product_Name', 'Category'])
for r, row in enumerate(SALES, 2):
    for c, v in enumerate(row, 1):
        cell = sd.cell(row=r, column=c, value=v); cell.fill = INPUT
        if c == 2: cell.number_format = 'yyyy-mm-dd'
    f(sd, f'G{r}', f'=E{r}*F{r}', row[4] * row[5], '#,##0')
    name = next((p[1] for p in PRODUCTS if p[0] == row[2]), 'ไม่พบสินค้า'); cat = next((p[2] for p in PRODUCTS if p[0] == row[2]), 'ไม่พบสินค้า')
    f(sd, f'H{r}', f'=IFERROR(INDEX(Product_Master!$B$2:$B$4,MATCH(C{r},Product_Master!$A$2:$A$4,0)),"ไม่พบสินค้า")', name)
    f(sd, f'I{r}', f'=IFERROR(VLOOKUP(C{r},Product_Master!$A$2:$C$4,3,FALSE),"ไม่พบสินค้า")', cat)
nS = len(SALES) + 1
sd[f'A{nS + 2}'] = 'ยอดขาย Online'; f(sd, f'G{nS + 2}', f'=SUMIF($D$2:$D${nS},"Online",$G$2:$G${nS})', sum(s[4] * s[5] for s in SALES if s[3] == 'Online'), '#,##0')
sd[f'A{nS + 3}'] = 'ยอดขายรวม'; f(sd, f'G{nS + 3}', f'=SUM($G$2:$G${nS})', sum(s[4] * s[5] for s in SALES), '#,##0')
sd[f'A{nS + 4}'] = 'รหัสสินค้าไม่พบใน Master'; f(sd, f'G{nS + 4}', f'=COUNTIF($H$2:$H${nS},"ไม่พบสินค้า")', sum(1 for s in SALES if s[2] not in [p[0] for p in PRODUCTS]))
sd[f'A{nS + 6}'] = 'แถว O001–O003 มาจากต้นฉบับข้อ 15 · O004–O006 เพิ่มเพื่อทดสอบ (สินค้าใหม่ และรหัสที่ไม่มีใน Master)'; sd[f'A{nS + 6}'].font = NOTE_F
widths(sd, [10, 12, 11, 9, 9, 10, 10, 14, 12])
pm = wb.create_sheet('Product_Master')
header(pm, 1, ['Product_ID', 'Product_Name', 'Category', 'Price'])
for r, p in enumerate(PRODUCTS, 2):
    for c, v in enumerate(p, 1): cell = pm.cell(row=r, column=c, value=v); cell.fill = INPUT
pm['A6'] = 'P001–P002 จากต้นฉบับข้อ 15 · P003 และคอลัมน์ Price เพิ่มเพื่อฝึก Lookup ราคา'; pm['A6'].font = NOTE_F
widths(pm, [11, 14, 12, 8])

# ---------------- Farm_Performance (section 15 verbatim) ----------------
ss = wb.create_sheet('Farm_Performance')
header(ss, 1, ['Farm_ID', 'Date', 'Age_Days', 'Body_Weight_g', 'Feed_Intake_kg', 'Mortality_Count', 'Birds_Placed'])
for r, row in enumerate(SPEC_SAMPLE, 2):
    for c, v in enumerate(row, 1):
        cell = ss.cell(row=r, column=c, value=v); cell.fill = INPUT
        if c == 2: cell.number_format = 'yyyy-mm-dd'
ss['A7'] = ('ข้อจำกัด (ต้นฉบับ): Feed Intake ในตัวอย่างไม่ได้ระบุชัดว่าเป็นรายวันหรือสะสม จึงห้ามคำนวณ FCR โดยไม่กำหนดนิยามและช่วงเวลาที่ตรงกันก่อน; '
            'Mortality Count ก็ต้องระบุว่าเป็นรายวันหรือสะสม — ชีต Raw_Data จึงใช้คอลัมน์ Feed_Cum_kg และ Mortality_Cum ที่ระบุว่าเป็น "สะสม"')
ss['A7'].font = NOTE_F; ss['A7'].alignment = Alignment(wrap_text=True); ss.merge_cells('A7:G9')
widths(ss, [10, 12, 10, 15, 15, 16, 13])

# ---------------- M365_Formulas (text, copy into Microsoft 365) ----------------
m3 = wb.create_sheet('M365_Formulas')
header(m3, 1, ['หัวข้อ', 'สูตร Microsoft 365 (คัดลอกไปวาง)', 'เวอร์ชัน', 'สูตรทางเลือก (Excel รุ่นเก่า)'])
M365 = [
    ('B03 ราคาจาก Master', '=XLOOKUP(A2,Master!$A$2:$A$100,Master!$C$2:$C$100,"ไม่พบสินค้า")', 'Excel 2021+ / M365 / Web', '=IFERROR(INDEX(Master!$C$2:$C$100,MATCH(A2,Master!$A$2:$A$100,0)),"ไม่พบสินค้า")'),
    ('B04 แยกรหัส BR-2026-001', '=TEXTSPLIT(A2,"-")', 'Excel 2024 / M365 / Web', '=LEFT(A2,FIND("-",A2)-1) · =MID(A2,FIND("-",A2)+1,4) · =RIGHT(A2,3)'),
    ('I01 รหัสซ้ำ + ไม่พบฟาร์ม', '=IF(COUNTIF(A2:A1000,G2)>1,"รหัสซ้ำ",XLOOKUP(G2,A2:A1000,D2:D1000,"ไม่พบฟาร์ม"))', 'Excel 2021+ / M365', '=IF(COUNTIF(A2:A1000,G2)>1,"รหัสซ้ำ",IFERROR(INDEX(D2:D1000,MATCH(G2,A2:A1000,0)),"ไม่พบฟาร์ม"))'),
    ('I04 รายการของฟาร์มที่เลือก เรียงใหม่→เก่า', '=SORTBY(FILTER(A2:E1000,A2:A1000=H2),FILTER(B2:B1000,A2:A1000=H2),-1)', 'Excel 2021+ / M365', 'ใช้ AutoFilter + Sort หรือ PivotTable'),
    ('I04 (จัดการกรณีไม่พบ)', '=IFERROR(SORTBY(FILTER(A2:E1000,A2:A1000=H2),FILTER(B2:B1000,A2:A1000=H2),-1),"ไม่พบข้อมูลของฟาร์มนี้")', 'Excel 2021+ / M365', '-'),
    ('I05 รวม Jan + Feb', '=VSTACK(Jan!A2:E100,Feb!A2:E100)', 'Excel 2024 / M365', 'Power Query: Append Queries'),
    ('A01 %CV ด้วย LET', '=LET(w,B2:B101,n,COUNT(w),avg,AVERAGE(w),IF(OR(n<2,avg=0),NA(),STDEV.S(w)/avg*100))', 'Excel 2021+ / M365', '=IF(OR(COUNT(B2:B101)<2,AVERAGE(B2:B101)=0),NA(),STDEV.S(B2:B101)/AVERAGE(B2:B101)*100)'),
    ('A02 CALC_FCR (Name Manager)', '=LAMBDA(feed_kg,weight_gain_kg,IF(AND(ISNUMBER(feed_kg),ISNUMBER(weight_gain_kg),weight_gain_kg>0),feed_kg/weight_gain_kg,NA()))', 'Excel 2024 / M365', '=IF(AND(ISNUMBER(A2),ISNUMBER(B2),B2>0),A2/B2,NA())'),
    ('A07 สรุป Feed รายฟาร์ม', '=GROUPBY(Raw_Data!A2:A25,Raw_Data!E2:E25,MAX)', 'Microsoft 365 (ตรวจช่องทางอัปเดต)', 'SUMIFS / PivotTable — ดูชีต Calculations ส่วนสรุปรายฟาร์ม'),
]
for r, row in enumerate(M365, 2):
    for c, v in enumerate(row, 1):
        cell = m3.cell(row=r, column=c, value=v); cell.alignment = Alignment(wrap_text=True, vertical='top')
        if c in (2, 4) and str(v).startswith('='): cell.data_type = 's'  # store as text, never evaluated here
m3['A12'] = 'สูตรในชีตนี้เก็บเป็นข้อความ (ไม่คำนวณ) เพราะต้องใช้ Microsoft 365 / Excel 2021+ — คัดลอกไปวางในเซลล์ว่างของไฟล์ที่ใช้ Excel รุ่นที่รองรับเพื่อทดสอบ'
m3['A12'].font = NOTE_F
widths(m3, [30, 70, 22, 60])

wb.calculation.fullCalcOnLoad = True  # Excel recalculates on open (file has no cached values)
os.makedirs(os.path.dirname(OUT), exist_ok=True)
wb.save(OUT)
with open(os.path.join(ROOT, 'scripts', 'excel', '.workbook-expected.json'), 'w', encoding='utf-8') as fh:
    json.dump(EXPECTED, fh, ensure_ascii=False, indent=0)
# shared sample data for the web dashboard (same numbers as the workbook)
sample = {
    'note': 'ข้อมูลสมมติเพื่อฝึก ไม่ใช่มาตรฐานสายพันธุ์ — สร้างโดย scripts/excel/make_workbook.py (ตรงกับชีต Raw_Data ใน sample-workbook.xlsx)',
    'chickWeightG': CHICK_G, 'reportDate': REPORT_DATE.isoformat(),
    'farms': [{'Farm_ID': x[0], 'Region': x[1], 'Houses': x[2], 'Birds_Placed': x[3], 'Placement_Date': x[4].isoformat(), 'Target_FCR_42d': x[5]} for x in FARMS],
    'rows': [{**r, 'Date': r['Date'].isoformat()} for r in ROWS],
    'weights': WEIGHTS,
}
with open(os.path.join(ROOT, 'content', 'excel', 'sample-data.json'), 'w', encoding='utf-8') as fh:
    json.dump(sample, fh, ensure_ascii=False, indent=1)
print(f'wrote {os.path.relpath(OUT, ROOT)} · {len(EXPECTED)} formula cells · rows {NRAW}')
